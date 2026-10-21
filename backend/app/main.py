from datetime import datetime, timezone
from pathlib import Path
import json, re, sqlite3, uuid
from io import BytesIO
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pypdf import PdfReader
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from xml.sax.saxutils import escape
from .rules import RULES, analyze

BASE = Path(__file__).resolve().parents[1]
DATA = BASE / "data"
UPLOADS = DATA / "uploads"
DB = DATA / "iscc.sqlite3"
MAX_BYTES = 20 * 1024 * 1024
app = FastAPI(title="ISCC Demo API", version="1.0")
app.add_middleware(CORSMiddleware, allow_origins=["http://127.0.0.1:5173", "http://localhost:5173"], allow_methods=["*"], allow_headers=["*"])

def connect():
    DATA.mkdir(exist_ok=True); UPLOADS.mkdir(exist_ok=True)
    c = sqlite3.connect(DB); c.row_factory = sqlite3.Row
    c.execute("CREATE TABLE IF NOT EXISTS audits (id TEXT PRIMARY KEY, name TEXT, created_at TEXT, score INTEGER, summary TEXT, counts TEXT, findings TEXT, chars INTEGER, pages INTEGER, stored_name TEXT)")
    c.commit(); return c

def record(row):
    return {**dict(row), "counts":json.loads(row["counts"]), "findings":json.loads(row["findings"])}

@app.get("/api/health")
def health(): return {"status":"ok","rules":len(RULES)}

@app.get("/api/rules")
def rules(): return [{"id":r["id"],"title":r["title"],"severity":r["severity"],"description":r["description"],"action":r["action"],"reference":"Internal demo rule"} for r in RULES]

@app.get("/api/audits")
def history():
    c=connect()
    try: return [record(r) for r in c.execute("SELECT * FROM audits ORDER BY created_at DESC")]
    finally: c.close()

@app.get("/api/audits/{audit_id}")
def get_audit(audit_id: str):
    c=connect()
    try:
        r=c.execute("SELECT * FROM audits WHERE id=?",(audit_id,)).fetchone()
        if not r: raise HTTPException(404,"Audit not found")
        return record(r)
    finally: c.close()

@app.post("/api/audits")
async def create_audit(file: UploadFile = File(...)):
    name = Path(file.filename or "document").name
    ext=Path(name).suffix.lower()
    if ext not in (".pdf",".docx"): raise HTTPException(400,"Choose a PDF or DOCX file.")
    raw=await file.read(MAX_BYTES+1)
    if not raw: raise HTTPException(400,"The selected file is empty.")
    if len(raw)>MAX_BYTES: raise HTTPException(413,"File exceeds the 20 MiB upload limit.")
    ident=uuid.uuid4().hex
    safe=re.sub(r"[^A-Za-z0-9._-]","_",name)[:120] or "document"; stored=f"{ident}_{safe}"
    try:
        if ext==".pdf":
            reader=PdfReader(BytesIO(raw)); text="\n".join(p.extract_text() or "" for p in reader.pages); pages=len(reader.pages)
        else:
            doc=Document(BytesIO(raw)); text="\n".join(p.text for p in doc.paragraphs); pages=0
            for table in doc.tables:
                text += "\n"+"\n".join(" ".join(cell.text for cell in row.cells) for row in table.rows)
    except Exception as e: raise HTTPException(400,f"Could not read this document: {e}")
    if not text.strip(): raise HTTPException(422,"No extractable text was found. Scanned PDFs need OCR, which this demo does not include.")
    score, findings, counts, summary=analyze(text)
    UPLOADS.mkdir(parents=True,exist_ok=True); (UPLOADS/stored).write_bytes(raw)
    created=datetime.now(timezone.utc).isoformat()
    c=connect()
    try:
        c.execute("INSERT INTO audits VALUES (?,?,?,?,?,?,?,?,?,?)",(ident,name,created,score,summary,json.dumps(counts),json.dumps(findings),len(text),pages,stored)); c.commit()
        return record(c.execute("SELECT * FROM audits WHERE id=?",(ident,)).fetchone())
    finally: c.close()

@app.get("/api/audits/{audit_id}/report.pdf")
def pdf_report(audit_id: str):
    audit=get_audit(audit_id); buffer=BytesIO(); styles=getSampleStyleSheet(); story=[]
    story += [Paragraph("ISCC — Industrial Safety Compliance Checker",styles["Title"]),Paragraph("Educational demo screening report",styles["Heading2"]),Spacer(1,10)]
    story += [Paragraph(f"Document: {escape(audit['name'])}",styles["BodyText"]),Paragraph(f"Analyzed: {escape(audit['created_at'])}",styles["BodyText"]),Paragraph(f"Demo score: {audit['score']}% ({7-sum(audit['counts'].values())}/7 demo checks matched)",styles["BodyText"]),Paragraph(escape(audit['summary']),styles["BodyText"]),Spacer(1,12)]
    for severity in ("CRITICAL","MAJOR","MINOR"):
        for f in audit["findings"]:
            if f["severity"]==severity:
                story += [Paragraph(f"{severity} — {escape(f['title'])}",styles["Heading3"]),Paragraph(f"{escape(f['description'])}<br/>Evidence: {escape(f['evidence'])}<br/>Reference: Internal demo rule<br/>Suggested action: {escape(f['remediation'])}",styles["BodyText"]),Spacer(1,8)]
    story += [Spacer(1,10),Paragraph("Disclaimer: This educational screening uses simple phrase matching. It is not AI, legal advice, comprehensive regulatory validation, or official certification. Results may be incomplete or incorrect and require qualified human review.",styles["Italic"])]
    SimpleDocTemplate(buffer,pagesize=letter,title="ISCC screening report").build(story); buffer.seek(0)
    return StreamingResponse(buffer,media_type="application/pdf",headers={"Content-Disposition":f'attachment; filename="ISCC_{audit_id[:8]}_report.pdf"'})
